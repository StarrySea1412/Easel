#!/usr/bin/env python3
"""pptx_build.py — Markdown 大纲 → PPTX + 幻灯页 HTML / 逐页 PNG / 整册 PDF。

一份 Markdown 大纲同时生成：
    deck.pptx    — python-pptx（内置主题：深色封面 + 浅色内容页），文本存字体名，渲染交给 PowerPoint
    deck.html    — 同一份规格的 16:9 整页幻灯 HTML（每页一个 .slide，离线可开）
    slide-NN.png — 每页一张截图（Playwright + Chromium）
    deck.pdf     — 整册 PDF（Chromium 打印，1280×720）

大纲格式（每行确定性解析，详见 SKILL.md）：
    # 主题               封面标题（必须有；其后、首个 ## 前的文本行=副标题）
    > 备注               当前页的演讲者备注（封面区=封面备注；进 pptx 的 notes，不进画面）
    ## 页标题            开始一页（要点页为默认版面）
    - 要点               条目
    ### 左栏标题         两栏版面：一页里出现 ≥2 个 ### 时启用，栏内 - 为该栏条目
    !42%                大数字版面：主数字（每页至多一个）
    普通段落             引用版面：段落为引文；以 —/-- 开头的行=署名

依赖：python-pptx（必需）；playwright + chromium（PNG/PDF；缺失或加 --skip-render 时跳过，
只产出 pptx/html 并给出安装指引）。

子命令：
    build     大纲 → 产物（默认输出目录内全部产物）
    selftest  自检（临时目录内全链路）

用法举例：
    pptx_build.py build -i outline.md -o outputs/年度复盘/
    pptx_build.py build -i outline.md -o outputs/年度复盘/ --skip-render
"""
from __future__ import annotations

import argparse
import html as _html
import sys
from dataclasses import dataclass, field
from pathlib import Path

FONT = "微软雅黑"
SLIDE_W = 12192000  # 13.333in，16:9
SLIDE_H = 6858000   # 7.5in

COVER_BG = "111827"
INK = "1F2937"
MUTED = "6B7280"
SUBTITLE = "9CA3AF"
ACCENT = "2563EB"
HL = "F59E0B"
PAGE_BG = "FFFFFF"


def _die(msg: str, code: int = 1) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(code)


# ---------------------------------------------------------------- 大纲解析

@dataclass
class Slide:
    kind: str = "bullets"          # bullets / two_column / big_number / quote
    title: str = ""
    bullets: list[str] = field(default_factory=list)
    columns: list[tuple[str, list[str]]] = field(default_factory=list)
    big_number: str = ""
    quote: str = ""
    attribution: str = ""
    notes: str = ""


@dataclass
class Deck:
    title: str
    subtitle: str = ""
    notes: str = ""                # 封面页演讲者备注
    slides: list[Slide] = field(default_factory=list)


def parse_outline(md_text: str) -> Deck:
    """Markdown 大纲 → Deck（确定性纯函数）。"""
    if not md_text or not md_text.strip():
        raise ValueError("大纲为空")
    deck: Deck | None = None
    current: Slide | None = None
    column: tuple[str, list[str]] | None = None
    paragraph_lines: list[str] = []
    subtitle_lines: list[str] = []

    def flush_paragraph() -> None:
        if current is None or not paragraph_lines:
            return
        text = " ".join(paragraph_lines).strip()
        paragraph_lines.clear()
        if not text:
            return
        if not current.quote:
            current.quote = text
        else:
            current.attribution = (current.attribution + " " + text).strip()

    def flush_slide() -> None:
        nonlocal current, column
        flush_paragraph()
        if current is None:
            return
        if current.bullets and not current.columns and not current.big_number and not current.quote:
            current.kind = "bullets"
        if current.columns and len(current.columns) >= 2 and not current.big_number and not current.quote:
            current.kind = "two_column"
        if current.big_number:
            current.kind = "big_number"
        if current.quote:
            current.kind = "quote"
        if deck is not None and (current.title or current.bullets or current.columns
                                 or current.big_number or current.quote or current.notes):
            deck.slides.append(current)
        current, column = None, None

    for raw in md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.strip()
        if not line:
            flush_paragraph()
            continue
        if line.startswith("# "):
            if deck is not None:
                raise ValueError("大纲只能有一个 # 主题行")
            flush_paragraph()
            deck = Deck(title=line[2:].strip())
            continue
        if deck is None:
            raise ValueError("大纲缺少 # 主题行（第一行应为 # 主题）")
        if line.startswith("## "):
            flush_slide()
            current = Slide(title=line[3:].strip())
            continue
        if line.startswith("### "):
            flush_paragraph()
            if current is None:
                raise ValueError("### 栏标题必须出现在某个 ## 页内")
            column = (line[4:].strip(), [])
            current.columns.append(column)
            continue
        if line.startswith("- ") or line.startswith("* "):
            flush_paragraph()
            item = line[2:].strip()
            if current is None:
                subtitle_lines.append(item)   # 封面区偶发要点：并入副标题，不新建页
            elif column is not None:
                column[1].append(item)
            else:
                current.bullets.append(item)
            continue
        if line.startswith("!") and current is not None:
            flush_paragraph()
            current.big_number = line[1:].strip()
            continue
        if line.startswith(">"):
            flush_paragraph()
            note = line[1:].strip()
            if current is None:
                deck.notes = (deck.notes + "\n" + note).strip()
            else:
                current.notes = (current.notes + "\n" + note).strip()
            continue
        # 普通段落：封面区=副标题；页内=引用文本；以 —/-- 开头=署名
        if current is None:
            subtitle_lines.append(line)
        elif line.startswith("—") or line.startswith("--"):
            flush_paragraph()
            current.attribution = (current.attribution + " " + line).strip()
        else:
            paragraph_lines.append(line)

    flush_slide()
    if deck is None:
        raise ValueError("大纲缺少 # 主题行（第一行应为 # 主题）")
    deck.subtitle = " ".join(subtitle_lines).strip() or deck.subtitle
    return deck


# ---------------------------------------------------------------- PPTX 生成

def _font(run, size: float, *, bold: bool = False, color: str = INK) -> None:
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = FONT  # latin 字体（python-pptx 内部按 schema 序插入 a:latin）
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    # CJK 需要显式 a:ea/a:cs 字体名，PowerPoint 才对中文生效
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def _para(tf, text: str, size: float, *, first: bool = False, bold: bool = False,
          color: str = INK, bullet: bool = False, align=None, space_after: float = 8,
          line_spacing: float = 1.15):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align or PP_ALIGN.LEFT
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = ("• " + text) if bullet else text
    _font(run, size, bold=bold, color=color)
    return p


def _textbox(slide, l, t, w, h):
    from pptx.util import Inches
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def _rect(slide, l, t, w, h, color: str):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _page_title(slide, title: str) -> None:
    tf = _textbox(slide, 0.9, 0.55, 11.5, 0.9)
    _para(tf, title, 28, first=True, bold=True, color=INK, space_after=0)
    _rect(slide, 0.92, 1.32, 0.72, 0.05, ACCENT)


def _new_slide(prs, bg: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    from pptx.dml.color import RGBColor
    slide.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    return slide


def write_pptx(deck: Deck, out_path: Path) -> Path:
    """Deck → pptx 文件（python-pptx 必需）。"""
    try:
        from pptx import Presentation
        from pptx.util import Emu, Pt  # noqa: F401
    except Exception as e:
        _die(f"需要 python-pptx：{e}（pip install python-pptx）", 3)
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)

    cover = _new_slide(prs, COVER_BG)
    tf = _textbox(cover, 0.9, 2.45, 11.5, 1.6)
    _para(tf, deck.title, 40, first=True, bold=True, color="FFFFFF", space_after=0)
    _rect(cover, 0.92, 3.72, 0.96, 0.06, HL)
    if deck.subtitle:
        stf = _textbox(cover, 0.9, 4.0, 11.5, 1.0)
        _para(stf, deck.subtitle, 17, first=True, color=SUBTITLE)
    if deck.notes:
        cover.notes_slide.notes_text_frame.text = deck.notes

    for s in deck.slides:
        slide = _new_slide(prs, PAGE_BG)
        _page_title(slide, s.title)
        if s.kind == "big_number":
            ntf = _textbox(slide, 0.9, 1.9, 11.5, 1.9)
            _para(ntf, s.big_number, 88, first=True, bold=True, color=ACCENT, space_after=0)
            if s.bullets:
                btf = _textbox(slide, 0.9, 4.2, 11.5, 2.6)
                for i, item in enumerate(s.bullets):
                    _para(btf, item, 16, first=(i == 0), bullet=True, color=MUTED)
        elif s.kind == "two_column":
            cols = s.columns[:2]
            for idx, (header, items) in enumerate(cols):
                left = 0.9 + idx * 6.0
                ctf = _textbox(slide, left, 1.7, 5.4, 5.1)
                _para(ctf, header, 20, first=True, bold=True, color=ACCENT, space_after=12)
                for item in items:
                    _para(ctf, item, 15, bullet=True)
        elif s.kind == "quote":
            qtf = _textbox(slide, 1.2, 2.5, 10.9, 2.2)
            _para(qtf, s.quote, 28, first=True, color=INK, space_after=0, line_spacing=1.3)
            if s.attribution:
                from pptx.enum.text import PP_ALIGN
                atf = _textbox(slide, 0.9, 4.9, 10.9, 0.7)
                _para(atf, s.attribution, 15, first=True, color=MUTED, align=PP_ALIGN.RIGHT)
        else:
            btf = _textbox(slide, 0.9, 1.75, 11.5, 5.0)
            for i, item in enumerate(s.bullets):
                _para(btf, item, 17, first=(i == 0), bullet=True, space_after=10)
        if s.notes:
            slide.notes_slide.notes_text_frame.text = s.notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


# ---------------------------------------------------------------- HTML 生成

_CSS = """
*{box-sizing:border-box;margin:0;padding:0;}
@page{size:1280px 720px;margin:0;}
body{background:#e5e7eb;font-family:'Microsoft YaHei','PingFang SC','Noto Sans CJK SC',sans-serif;}
.slide{width:1280px;height:720px;position:relative;overflow:hidden;background:#fff;
  page-break-after:always;break-after:page;margin:0 auto;}
.slide.cover{background:#111827;}
.accent-bar{position:absolute;background:#2563eb;}
.page-title{position:absolute;left:86px;top:53px;font-size:37px;font-weight:700;color:#1f2937;}
.rule{position:absolute;left:88px;top:127px;width:69px;height:5px;background:#2563eb;}
.bullets{position:absolute;left:86px;top:168px;width:1104px;}
.bullets div{font-size:23px;color:#1f2937;margin-bottom:14px;line-height:1.35;}
.bullets div:before{content:"• ";color:#2563eb;}
.big{position:absolute;left:86px;top:182px;font-size:117px;font-weight:700;color:#2563eb;line-height:1.05;}
.big-expl{position:absolute;left:86px;top:404px;width:1104px;}
.big-expl div{font-size:21px;color:#6b7280;margin-bottom:12px;line-height:1.35;}
.big-expl div:before{content:"• ";color:#2563eb;}
.cols{position:absolute;left:86px;top:163px;width:1104px;display:flex;gap:54px;}
.col{width:525px;}
.col h3{font-size:27px;font-weight:700;color:#2563eb;margin-bottom:18px;}
.col div{font-size:20px;color:#1f2937;margin-bottom:12px;line-height:1.4;}
.col div:before{content:"• ";color:#2563eb;}
.quote-text{position:absolute;left:115px;top:240px;width:1050px;font-size:37px;color:#1f2937;line-height:1.3;}
.attrib{position:absolute;left:86px;right:86px;top:470px;text-align:right;font-size:20px;color:#6b7280;}
.cover-title{position:absolute;left:86px;top:235px;width:1104px;font-size:53px;font-weight:700;color:#fff;}
.cover-bar{position:absolute;left:88px;top:357px;width:92px;height:6px;background:#f59e0b;}
.cover-sub{position:absolute;left:86px;top:384px;width:1104px;font-size:23px;color:#9ca3af;}
"""

_HTML = """<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"><title>{title}</title>
<style>{css}</style></head><body>{slides}</body></html>"""


def _esc(text: str) -> str:
    return _html.escape(text, quote=True)


def render_html(deck: Deck) -> str:
    """Deck → 整页幻灯 HTML（纯函数，供测试）。"""
    parts: list[str] = []
    sub = f'<div class="cover-sub">{_esc(deck.subtitle)}</div>' if deck.subtitle else ""
    parts.append('<section class="slide cover">'
                 f'<div class="cover-title">{_esc(deck.title)}</div>'
                 '<div class="cover-bar"></div>' + sub + "</section>")
    for s in deck.slides:
        body = ""
        if s.kind == "big_number":
            body += f'<div class="big">{_esc(s.big_number)}</div>'
            if s.bullets:
                items = "".join(f"<div>{_esc(b)}</div>" for b in s.bullets)
                body += f'<div class="big-expl">{items}</div>'
        elif s.kind == "two_column":
            cols = "".join(
                f'<div class="col"><h3>{_esc(h)}</h3>'
                + "".join(f"<div>{_esc(i)}</div>" for i in items) + "</div>"
                for h, items in s.columns[:2])
            body += f'<div class="cols">{cols}</div>'
        elif s.kind == "quote":
            body += f'<div class="quote-text">{_esc(s.quote)}</div>'
            if s.attribution:
                body += f'<div class="attrib">{_esc(s.attribution)}</div>'
        else:
            items = "".join(f"<div>{_esc(b)}</div>" for b in s.bullets)
            body += f'<div class="bullets">{items}</div>'
        parts.append(f'<section class="slide"><div class="page-title">{_esc(s.title)}</div>'
                     '<div class="rule"></div>' + body + "</section>")
    return _HTML.format(title=_esc(deck.title), css=_CSS, slides="".join(parts))


def render_images(html_path: Path, out_dir: Path, base: str) -> dict:
    """HTML → slide-NN.png + <base>.pdf；依赖缺失时返回 {rendered: False, reason}。"""
    try:
        from playwright.sync_api import sync_playwright
    except Exception as e:
        return {"rendered": False, "reason": f"需要 playwright：{e}（pip install playwright && playwright install chromium）"}
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": 1280, "height": 720}, device_scale_factor=2)
            page.goto(html_path.as_uri(), wait_until="networkidle", timeout=30000)
            slides = page.locator(".slide")
            count = slides.count()
            for i in range(count):
                slides.nth(i).screenshot(path=str(out_dir / f"slide-{i + 1:02d}.png"))
            page.pdf(path=str(out_dir / f"{base}.pdf"), width="1280px", height="720px",
                     print_background=True)
            browser.close()
            return {"rendered": True, "slides": count}
    except Exception as e:
        return {"rendered": False, "reason": f"Chromium 渲染失败：{e}（可 playwright install chromium 后重试）"}


# ---------------------------------------------------------------- CLI

def cmd_build(a) -> int:
    src = Path(a.input).expanduser()
    if not src.is_file():
        _die(f"输入不存在：{src}")
    try:
        deck = parse_outline(src.read_text(encoding="utf-8"))
    except ValueError as e:
        _die(f"大纲解析失败：{e}")
    out_dir = Path(a.output).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    pptx_path = write_pptx(deck, out_dir / f"{a.name}.pptx")
    print(f"✅ {pptx_path} ({pptx_path.stat().st_size / 1024:.0f} KB) — {1 + len(deck.slides)} 页")

    html_text = render_html(deck)
    html_path = out_dir / f"{a.name}.html"
    html_path.write_text(html_text, encoding="utf-8")
    print(f"✅ {html_path} ({html_path.stat().st_size / 1024:.0f} KB) HTML")

    if a.skip_render:
        print("⏭️  已按 --skip-render 跳过 PNG/PDF 渲染（pptx/html 不受影响）")
        return 0
    result = render_images(html_path, out_dir, a.name)
    if not result["rendered"]:
        print(f"⚠️  PNG/PDF 预览已跳过：{result['reason']}", file=sys.stderr)
        return 0
    count = result["slides"]
    print(f"✅ slide-01..{count:02d}.png（每页一张）+ {a.name}.pdf")
    return 0


def _sample_outline() -> str:
    return (
        "# 年度内容复盘\n"
        "2026 上半年 · 创作团队内部汇报\n"
        "\n"
        "> 开场：先说结论，再讲证据。\n"
        "\n"
        "## 三个关键数字\n"
        "!42%\n"
        "- 收藏率相对上季度的提升幅度\n"
        "- 主要来自「流程类」选题\n"
        "> 这个数字先核对数据源再讲。\n"
        "\n"
        "## 去年的做法 vs 今年\n"
        "### 去年\n"
        "- 单平台发布\n"
        "- 每周 3 条，凭手感选题\n"
        "### 今年\n"
        "- 多平台并进，选题先过结构模板\n"
        "- 每周 2 条，但收藏率翻倍\n"
        "\n"
        "## 一句共识\n"
        "“结构化比灵感更可靠。”\n"
        "— 团队复盘会\n"
    )


def cmd_selftest(_a) -> int:
    import tempfile
    deck = parse_outline(_sample_outline())
    assert deck.title == "年度内容复盘", "封面标题解析错误"
    assert deck.subtitle.startswith("2026 上半年"), "副标题解析错误"
    assert [s.kind for s in deck.slides] == ["big_number", "two_column", "quote"], \
        f"版面识别错误：{[s.kind for s in deck.slides]}"
    assert deck.slides[0].big_number == "42%" and len(deck.slides[0].bullets) == 2
    assert deck.notes.startswith("开场"), "封面备注解析错误"
    assert deck.slides[0].notes.startswith("这个数字"), "页面备注未挂到本页"
    assert len(deck.slides[1].columns) == 2 and deck.slides[1].columns[1][1][0].startswith("多平台")
    assert deck.slides[2].attribution.startswith("—"), "署名解析错误"

    html_text = render_html(deck)
    assert html_text.count('class="slide') == 4, "HTML 页数错误"
    assert "年度内容复盘" in html_text and "42%" in html_text

    with tempfile.TemporaryDirectory() as td:
        out = Path(td)
        pptx_path = write_pptx(deck, out / "deck.pptx")
        assert pptx_path.is_file() and pptx_path.stat().st_size > 5000, "pptx 未生成"
        from pptx import Presentation
        prs = Presentation(str(pptx_path))
        assert len(prs.slides) == 4, f"pptx 页数错误：{len(prs.slides)}"
        assert prs.slide_width == SLIDE_W and prs.slide_height == SLIDE_H, "幻灯片尺寸不是 16:9"
        texts = []
        for sl in prs.slides:
            for shape in sl.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        joined = "\n".join(texts)
        assert "年度内容复盘" in joined and "42%" in joined and "多平台并进" in joined, "pptx 文本缺失"
        assert prs.slides[0].notes_slide.notes_text_frame.text.startswith("开场"), "封面备注缺失"
        assert prs.slides[1].notes_slide.notes_text_frame.text.startswith("这个数字"), "pptx 页面备注缺失"

        html_path = out / "deck.html"
        html_path.write_text(render_html(deck), encoding="utf-8")
        result = render_images(html_path, out, "deck")
        if not result["rendered"]:
            print(f"↷ 渲染自检跳过：{result['reason']}", file=sys.stderr)
        else:
            assert (out / "slide-01.png").is_file() and (out / "slide-04.png").is_file(), "逐页 PNG 缺失"
            assert (out / "deck.pdf").is_file() and (out / "deck.pdf").stat().st_size > 5000, "PDF 缺失"
            print("✅ selftest 通过（大纲→pptx 读回校验→HTML→逐页 PNG + PDF）")
            return 0
    print("✅ selftest 通过（大纲→pptx 读回校验→HTML；PNG/PDF 因缺依赖跳过）")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Markdown 大纲 → PPTX + 幻灯页 HTML/PNG/PDF（python-pptx + Chromium）",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd")

    p = sub.add_parser("build", help="大纲 → 产物")
    p.add_argument("-i", "--input", required=True, help="Markdown 大纲文件")
    p.add_argument("-o", "--output", required=True, help="输出目录")
    p.add_argument("--name", default="deck", help="产物基础名（默认 deck）")
    p.add_argument("--skip-render", action="store_true", help="只出 pptx/html，跳过 PNG/PDF")
    p.set_defaults(func=cmd_build)

    sub.add_parser("selftest", help="自检").set_defaults(func=cmd_selftest)

    a = ap.parse_args()
    if not getattr(a, "func", None):
        ap.print_help()
        return 1
    return a.func(a)


if __name__ == "__main__":
    sys.exit(main())
