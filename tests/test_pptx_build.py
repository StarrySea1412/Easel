"""pptx-slides 生成器（skills/shared/scripts/pptx_build.py）：解析、生成、CLI 与渲染。"""
import subprocess
import os
import sys
from pathlib import Path

import pytest

SCRIPT = Path(__file__).resolve().parents[1] / 'skills' / 'shared' / 'scripts' / 'pptx_build.py'
sys.path.insert(0, str(SCRIPT.parent))
import pptx_build  # noqa: E402


OUTLINE = """\
# 年度内容复盘
2026 上半年 · 团队内部

> 开场：先说结论。

## 三个关键数字
!42%
- 收藏率提升
- 来自流程类选题
> 先核对数据源。

## 去年 vs 今年
### 去年
- 单平台
- 凭手感选题
### 今年
- 多平台
- 结构模板立项

## 一句共识
“结构化比灵感更可靠。”
— 团队复盘会
"""


def test_parse_outline_detects_cover_pages_and_notes():
    deck = pptx_build.parse_outline(OUTLINE)
    assert deck.title == '年度内容复盘'
    assert deck.subtitle == '2026 上半年 · 团队内部'
    assert deck.notes == '开场：先说结论。'
    assert [s.kind for s in deck.slides] == ['big_number', 'two_column', 'quote']
    assert deck.slides[0].big_number == '42%' and len(deck.slides[0].bullets) == 2
    assert deck.slides[0].notes == '先核对数据源。'
    assert deck.slides[1].columns[0] == ('去年', ['单平台', '凭手感选题'])
    assert deck.slides[1].columns[1][1] == ['多平台', '结构模板立项']
    assert deck.slides[2].quote == '“结构化比灵感更可靠。”'
    assert deck.slides[2].attribution == '— 团队复盘会'


def test_default_page_kind_is_bullets():
    deck = pptx_build.parse_outline('# T\n\n## 页\n- a\n- b\n')
    assert deck.slides[0].kind == 'bullets' and deck.slides[0].bullets == ['a', 'b']


@pytest.mark.parametrize('bad', [
    '', '   ', '无井号\n## 页\n- a', '# 一\n# 二\n', '# T\n### 栏\n- a\n',
])
def test_parse_rejects_invalid_outlines(bad):
    with pytest.raises(ValueError):
        pptx_build.parse_outline(bad)


def test_write_pptx_round_trips_pages_text_size_and_notes(tmp_path):
    from pptx.oxml.ns import qn
    from pptx import Presentation
    deck = pptx_build.parse_outline(OUTLINE)
    out = pptx_build.write_pptx(deck, tmp_path / 'deck.pptx')
    assert out.is_file() and out.stat().st_size > 5000
    prs = Presentation(str(out))
    assert len(prs.slides) == 4, '封面 + 3 页'
    assert prs.slide_width == pptx_build.SLIDE_W and prs.slide_height == pptx_build.SLIDE_H
    joined = '\n'.join(sh.text_frame.text for sl in prs.slides
                       for sh in sl.shapes if sh.has_text_frame)
    for needle in ('年度内容复盘', '2026 上半年', '42%', '收藏率提升', '去年',
                   '多平台', '结构化比灵感更可靠', '— 团队复盘会'):
        assert needle in joined, f'pptx 缺少文本：{needle}'
    assert prs.slides[0].notes_slide.notes_text_frame.text == '开场：先说结论。'
    assert prs.slides[1].notes_slide.notes_text_frame.text == '先核对数据源。'
    title_shape = next(sh for sh in prs.slides[1].shapes
                       if sh.has_text_frame and '三个关键数字' in sh.text_frame.text)
    rPr = title_shape.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    assert rPr.find(qn('a:ea')).get('typeface') == pptx_build.FONT, 'CJK 字体名未写入 ea'


def test_html_renders_slides_and_escapes_user_text():
    deck = pptx_build.parse_outline('# T & <主题>\n\n## 页 <一>\n- a<b\n')
    html_text = pptx_build.render_html(deck)
    assert html_text.count('class="slide') == 2
    assert 'T &amp; &lt;主题&gt;' in html_text
    assert 'a&lt;b' in html_text
    assert 'width:1280px;height:720px' in html_text, '缺 16:9 页尺寸'


def test_build_cli_end_to_end_without_render(tmp_path):
    src = tmp_path / 'outline.md'
    src.write_text(OUTLINE, encoding='utf-8')
    out_dir = tmp_path / 'out'
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), 'build', '-i', str(src), '-o', str(out_dir), '--skip-render'],
        capture_output=True, text=True, encoding='utf-8', env={**os.environ, 'PYTHONUTF8': '1'})
    assert proc.returncode == 0, proc.stderr
    assert (out_dir / 'deck.pptx').is_file() and (out_dir / 'deck.html').is_file()
    assert '已按 --skip-render 跳过' in proc.stdout


def test_build_cli_rejects_bad_outline(tmp_path):
    src = tmp_path / 'bad.md'
    src.write_text('没有主题行\n', encoding='utf-8')
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), 'build', '-i', str(src), '-o', str(tmp_path / 'o')],
        capture_output=True, text=True, encoding='utf-8', env={**os.environ, 'PYTHONUTF8': '1'})
    assert proc.returncode != 0
    assert 'ERROR' in proc.stderr and '大纲' in proc.stderr


def test_render_png_and_pdf_or_explicit_skip(tmp_path):
    deck = pptx_build.parse_outline(OUTLINE)
    html_path = tmp_path / 'deck.html'
    html_path.write_text(pptx_build.render_html(deck), encoding='utf-8')
    result = pptx_build.render_images(html_path, tmp_path, 'deck')
    if not result['rendered']:
        pytest.skip(f"环境无 Playwright/Chromium：{result['reason']}")
    assert result['slides'] == 4
    for n in range(1, 5):
        png = tmp_path / f'slide-{n:02d}.png'
        assert png.is_file() and png.stat().st_size > 10000, f'{png.name} 缺失或过小'
    assert (tmp_path / 'deck.pdf').stat().st_size > 5000
